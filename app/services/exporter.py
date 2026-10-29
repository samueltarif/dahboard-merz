"""Exportação para Excel (várias abas formatadas) e CSV.

A aba "SAP" sai com os cabeçalhos originais do SAP, então o arquivo exportado
pode ser importado de novo sem ajustes.
"""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.columns import COLUMNS
from app.domain.taxes import TAXES
from app.services import compliance, metrics, reform
from app.services.filters import Filters

HEADER_FILL = PatternFill("solid", fgColor="1F1E1C")
HEADER_FONT = Font(bold=True, color="F7F6F3")
TITLE_FONT = Font(bold=True, size=14)
MONEY_FMT = '"R$" #,##0.00'
PCT_FMT = "0.00%"

DERIVED = {"fluxo": "Movimento", "operacao": "Operação", "natureza": "Natureza", "uf_op": "UF operação", "periodo": "Período", "cfop": "CFOP (4 dígitos)"}

LABELS = {
    "fluxo": "Movimento", "operacao": "Operação", "natureza": "Natureza", "uf_op": "UF", "cfop": "CFOP",
    "natureza_operacao": "Descrição", "periodo": "Período", "dia": "Dia", "documentos": "Documentos", "itens": "Itens",
    "quantidade": "Quantidade", "valor_total": "Valor total c/ impostos", "valor": "Valor produtos", "net_sales": "NET Sales",
    "tributos": "Tributos (consumo)", "tributos_reforma": "IBS+CBS", "carga": "Carga efetiva", "material": "Material",
    "material_desc": "Descrição", "ncm": "NCM", "parceiro_id": "ID parceiro", "parceiro_nome": "Parceiro", "preco_medio": "Preço médio",
    **{t.key: t.label for t in TAXES},
}
PCT_COLS = {"Carga efetiva", "Participação", "Carga"}
NO_MONEY = {"Documentos", "Itens", "Quantidade", "CFOP", "NCM", "Material", "ID parceiro", "Período", "Dia", "UF", "Alíq.", "Regras"}


def _frame(records: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(records)
    return df.rename(columns={c: LABELS.get(c, c) for c in df.columns})


def _style(ws, df: pd.DataFrame, start_row: int = 1) -> None:
    for j, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=j)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200).tolist())) if len(df) else len(str(col))
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 48)
        numeric = pd.api.types.is_numeric_dtype(df[col])
        if not numeric:
            continue
        fmt = PCT_FMT if col in PCT_COLS else (None if col in NO_MONEY or "Alíq" in str(col) or "%" in str(col) else MONEY_FMT)
        if fmt:
            for i in range(start_row + 1, start_row + 1 + len(df)):
                ws.cell(row=i, column=j).number_format = fmt
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if len(df):
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}"


def _write(writer: pd.ExcelWriter, name: str, df: pd.DataFrame) -> None:
    df.to_excel(writer, sheet_name=name, index=False)
    _style(writer.sheets[name], df)


def _sap_layout(df: pd.DataFrame) -> pd.DataFrame:
    """Base filtrada com cabeçalhos do SAP + colunas derivadas no fim."""
    out = pd.DataFrame({c.sap: df[c.field] for c in COLUMNS})
    for c in ("data_documento", "data_lancamento"):
        out[next(x.sap for x in COLUMNS if x.field == c)] = pd.to_datetime(df[c]).dt.date
    for f, label in DERIVED.items():
        out[label] = df[f]
    return out


def _filters_text(f: Filters) -> list[tuple[str, str]]:
    def j(v):
        return ", ".join(map(str, v)) if v else "Todos"
    return [
        ("Movimento", f.fluxo or "Todos"), ("Operação", j(f.operacao)), ("Natureza", j(f.natureza)), ("UF", j(f.uf)),
        ("Modelo", j(f.modelo)), ("Lotes", j(f.imports)),
        ("Período", f"{f.periodo_ini or 'início'} a {f.periodo_fim or 'fim'}"), ("Busca", f.busca or "—"),
    ]


def to_excel(df: pd.DataFrame, df_sem_fluxo: pd.DataFrame, df_apuracao: pd.DataFrame, filters: Filters) -> bytes:
    ov = metrics.overview(df, df_sem_fluxo, df_apuracao)
    pr = metrics.products(df, limit=10_000)
    rf = reform.simulate(df)
    cp = compliance.run(df, sample=100_000)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        # --- Resumo -------------------------------------------------------------
        k = ov["kpis"]
        resumo = pd.DataFrame(
            [("Valor total c/ impostos", k["valor_total"]["value"]), ("Valor dos produtos", k["valor_produtos"]),
             ("NET Sales", k["net_sales"]), ("Tributos sobre consumo", k["tributos"]["value"]),
             ("Carga efetiva", k["carga"]["value"]), ("IBS + CBS (teste 2026)", k["tributos_reforma"]),
             ("Documentos", k["documentos"]["value"]), ("Itens", k["itens"]), ("Parceiros", k["parceiros"])],
            columns=["Indicador", "Valor"],
        )
        resumo.to_excel(w, sheet_name="Resumo", index=False, startrow=11)
        ws = w.sheets["Resumo"]
        ws["A1"], ws["A1"].font = "Painel tributário — exportação", TITLE_FONT
        for i, (lab, val) in enumerate(_filters_text(filters), start=3):
            ws.cell(row=i, column=1, value=lab).font = Font(bold=True)
            ws.cell(row=i, column=2, value=val)
        _style(ws, resumo, start_row=12)
        for i, (lab, _) in enumerate(resumo.values.tolist(), start=13):
            ws.cell(row=i, column=2).number_format = PCT_FMT if lab == "Carga efetiva" else ("0" if lab in ("Documentos", "Itens", "Parceiros") else MONEY_FMT)
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 26

        tx = pd.DataFrame(ov["taxes"])[["label", "value", "share", "base_carga"]]
        tx.columns = ["Tributo", "Valor", "Participação", "Carga"]
        _write(w, "Tributos", tx)
        ap = pd.DataFrame(ov["apuracao"]).drop(columns=["key"], errors="ignore")
        ap.columns = ["Tributo", "Débitos (saídas)", "Créditos (entradas)", "Sem direito a crédito", "Saldo"][: len(ap.columns)]
        _write(w, "Débito x crédito", ap)
        _write(w, "Por operação", _frame(ov["by_operacao"]))
        _write(w, "Por natureza", _frame(ov["by_natureza"]))
        _write(w, "Por UF", _frame(ov["by_uf"]))
        _write(w, "Por CFOP", _frame(ov["by_cfop"]))
        _write(w, "Por dia", _frame(ov["timeline"]))
        _write(w, "Materiais", _frame(pr["materiais"]))
        _write(w, "Parceiros", _frame(pr["parceiros"]))
        _write(w, "NCM", _frame(pr["ncm"]))
        if rf.get("por_cst"):
            _write(w, "Reforma IBS-CBS", pd.DataFrame(rf["por_cst"]).rename(columns={
                "cst_desc": "CST IBS/CBS", "itens": "Itens", "valor_total": "Valor total c/ impostos", "base": "Base IBS/CBS",
                "ibs_cbs_2026": "IBS+CBS destacado 2026", "atual": "Tributos substituídos (atual)", "simulado": "IBS+CBS simulado"}))
        cps = pd.DataFrame([{"Regra": c["title"], "Severidade": c["severity"], "Itens": c["count"], "Documentos": c["documentos"],
                             "Valor envolvido": c["valor"], "Descrição": c["detail"]} for c in cp["checks"]])
        _write(w, "Conformidade", cps)
        _write(w, "SAP", _sap_layout(df))
    return buf.getvalue()


def to_csv(df: pd.DataFrame) -> bytes:
    """CSV no padrão do Excel brasileiro: ';' como separador, vírgula decimal, datas dd/mm/aaaa."""
    out = _sap_layout(df)
    for c in ("data_documento", "data_lancamento"):
        header = next(x.sap for x in COLUMNS if x.field == c)
        out[header] = pd.to_datetime(df[c]).dt.strftime("%d/%m/%Y")
    return out.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig")
