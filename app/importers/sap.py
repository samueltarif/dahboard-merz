"""Leitura de planilhas com o relatório fiscal do SAP.

Fluxo:  bytes do arquivo -> escolhe a aba -> acha a linha de cabeçalho ->
        renomeia colunas (app.domain.columns) -> converte tipos -> DataFrame.

Aceita .xlsx/.xlsm e .csv (separador ; ou , e números no formato brasileiro).
Nada aqui toca no banco: a função é pura e fácil de testar.
"""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.config import settings
from app.domain.columns import COLUMNS, DATE_FIELDS, NUMERIC_FIELDS, REQUIRED_HEADERS, TEXT_FIELDS
from app.importers.errors import ImportFileError

EXCEL_EXT = {".xlsx", ".xlsm"}
CSV_EXT = {".csv", ".txt"}
HEADER_SCAN_ROWS = 15


@dataclass
class ParsedSheet:
    filename: str
    sheet: str
    frame: pd.DataFrame          # já com nomes internos e tipos convertidos
    unknown_columns: list[str]   # colunas da planilha que não estão no mapa (ignoradas)


def _norm(text: object) -> str:
    """Normaliza cabeçalho para comparação: sem acento, minúsculo, sem espaços extras."""
    s = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


_HEADER_TO_FIELD: dict[str, str] = {}
for _c in COLUMNS:
    for _h in _c.headers:
        _HEADER_TO_FIELD.setdefault(_norm(_h), _c.field)
_REQUIRED_NORM = {_norm(h) for h in REQUIRED_HEADERS}


# ---------------------------------------------------------------------------
# Escolha da aba e do cabeçalho
# ---------------------------------------------------------------------------
def _header_score(values: list[object]) -> int:
    return len(_REQUIRED_NORM & {_norm(v) for v in values if v is not None})


def _pick_excel_sheet(content: bytes) -> tuple[str, int]:
    """Retorna (nome_da_aba, índice_da_linha_de_cabeçalho)."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        preferred = [n for n in names if n.strip() in settings.sap_sheet_names]
        ordered = preferred + [n for n in names if n not in preferred]
        for name in ordered:
            ws = wb[name]
            for idx, row in enumerate(ws.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True)):
                if _header_score(list(row)) == len(_REQUIRED_NORM):
                    return name, idx
        raise ImportFileError(
            "Nenhuma aba com o layout do SAP foi encontrada. "
            f"Abas no arquivo: {', '.join(names)}. "
            f"Colunas obrigatórias: {', '.join(REQUIRED_HEADERS)}.",
        )
    finally:
        wb.close()


def _read_excel(content: bytes) -> tuple[str, pd.DataFrame]:
    sheet, header_idx = _pick_excel_sheet(content)
    frame = pd.read_excel(io.BytesIO(content), sheet_name=sheet, header=header_idx, dtype=object, engine="openpyxl")
    return sheet, frame


def _read_csv(content: bytes) -> tuple[str, pd.DataFrame]:
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        raise ImportFileError("Não foi possível ler o CSV (codificação desconhecida).")
    first_line = text.splitlines()[0] if text else ""
    sep = ";" if first_line.count(";") >= first_line.count(",") else ","
    frame = pd.read_csv(io.StringIO(text), sep=sep, dtype=object)
    if _header_score(list(frame.columns)) < len(_REQUIRED_NORM):
        raise ImportFileError(f"O CSV não tem as colunas obrigatórias: {', '.join(REQUIRED_HEADERS)}.")
    return "csv", frame


# ---------------------------------------------------------------------------
# Conversão de tipos
# ---------------------------------------------------------------------------
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}")
_BR_NUMBER = re.compile(r"^-?\d{1,3}(\.\d{3})*(,\d+)?$|^-?\d+,\d+$")


def _to_number(series: pd.Series) -> pd.Series:
    def conv(v: object) -> object:
        if isinstance(v, str):
            s = v.strip().replace("R$", "").replace(" ", "")
            if not s:
                return None
            if _BR_NUMBER.match(s):
                s = s.replace(".", "").replace(",", ".")
            return s
        return v

    return pd.to_numeric(series.map(conv), errors="coerce").fillna(0.0).astype(float)


def _to_text(series: pd.Series) -> pd.Series:
    def conv(v: object) -> object:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return None
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        s = str(v).strip()
        return s or None

    return series.map(conv).astype(object)


def _to_date(series: pd.Series) -> pd.Series:
    def conv(v: object) -> object:
        if v is None or (isinstance(v, float) and pd.isna(v)) or v == "":
            return None
        if isinstance(v, str):
            s = v.strip()
            if _ISO_DATE.match(s):          # 2026-09-01 (ano primeiro)
                ts = pd.to_datetime(s[:10], format="%Y-%m-%d", errors="coerce")
            else:                           # 01/09/2026 (padrão brasileiro)
                ts = pd.to_datetime(s, dayfirst=True, errors="coerce")
        else:
            ts = pd.to_datetime(v, errors="coerce")
        return None if pd.isna(ts) else ts.date()

    return series.map(conv).astype(object)


def normalize_frame(raw: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Renomeia colunas SAP -> campos internos e converte os tipos."""
    rename: dict[str, str] = {}
    unknown: list[str] = []
    used: set[str] = set()
    for col in raw.columns:
        field = _HEADER_TO_FIELD.get(_norm(col))
        if field and field not in used:
            rename[col] = field
            used.add(field)
        else:
            unknown.append(str(col))

    frame = raw[list(rename)].rename(columns=rename).copy()
    for field in (c.field for c in COLUMNS):
        if field not in frame.columns:
            frame[field] = None

    for f in TEXT_FIELDS:
        frame[f] = _to_text(frame[f])
    for f in NUMERIC_FIELDS:
        frame[f] = _to_number(frame[f])
    for f in DATE_FIELDS:
        frame[f] = _to_date(frame[f])

    # Remove linhas vazias (totais, rodapés) e itens estornados.
    frame = frame[frame["numero_documento"].notna() & frame["cfop_raw"].notna()]
    frame = frame[frame["estornado"].fillna("").str.upper().ne("X")]
    return frame[[c.field for c in COLUMNS]].reset_index(drop=True), unknown


def parse_file(filename: str, content: bytes) -> ParsedSheet:
    ext = Path(filename).suffix.lower()
    if not content:
        raise ImportFileError("Arquivo vazio.")
    if ext in EXCEL_EXT:
        sheet, raw = _read_excel(content)
    elif ext in CSV_EXT:
        sheet, raw = _read_csv(content)
    elif ext == ".xls":
        raise ImportFileError("Formato .xls (Excel 97-2003) não suportado. Salve como .xlsx e importe de novo.")
    else:
        raise ImportFileError(f"Extensão '{ext or '?'}' não suportada. Use .xlsx, .xlsm ou .csv.")

    frame, unknown = normalize_frame(raw)
    if frame.empty:
        raise ImportFileError(f"A aba '{sheet}' não tem linhas válidas.")
    return ParsedSheet(filename=filename, sheet=sheet, frame=frame, unknown_columns=unknown)
